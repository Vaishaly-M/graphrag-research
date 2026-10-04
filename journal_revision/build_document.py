"""Build the journal revision Word document from the review's Markdown source."""
from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SOURCE = HERE / 'GraphRAG_Experiment_Revision_Plan.md'


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for piece in re.split(pattern, text):
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', piece)
        if match:
            label, url = match.groups()
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            style = OxmlElement('w:rStyle')
            style.set(qn('w:val'), 'Hyperlink')
            props.append(style)
            run.append(props)
            content = OxmlElement('w:t')
            content.text = label
            run.append(content)
            link.append(run)
            paragraph._p.append(link)
        elif piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif piece.startswith('**') and piece.endswith('**'):
            paragraph.add_run(piece[2:-2]).bold = True
        else:
            paragraph.add_run(piece)


def add_table(doc, rows):
    # The report's table content does not contain escaped literal pipes.
    cells = [[cell.strip() for cell in row.strip().strip('|').split('|')] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r':?-+:?', cell or ' ') for cell in row)]
    table = doc.add_table(rows=1, cols=len(cells[0]))
    table.style = 'Table Grid'
    table.autofit = True
    for index, values in enumerate(cells):
        row = table.rows[0] if index == 0 else table.add_row()
        for cell, value in zip(row.cells, values):
            inline(cell.paragraphs[0], value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9)
                    if index == 0:
                        run.bold = True
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E6EEF5')
                cell._tc.get_or_add_tcPr().append(shading)
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    source = SOURCE.read_text(encoding='utf-8')
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('17365D')
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = 'GraphRAG | Journal experiment revision plan'
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string('666666')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('3 October 2026  |  Page ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith('# '):
            inline(doc.add_paragraph(style='Title'), line[2:])
        elif line.startswith('## '):
            inline(doc.add_paragraph(style='Heading 1'), line[3:])
        elif line.startswith('### '):
            inline(doc.add_paragraph(style='Heading 2'), line[4:])
        elif line.startswith('- [ ] '):
            inline(doc.add_paragraph(style='List Bullet'), '[ ] ' + line[6:])
        elif line.startswith('- '):
            inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            inline(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
        elif line.strip():
            inline(doc.add_paragraph(), line)
        i += 1
    doc.core_properties.title = 'GraphRAG Experiment Revision Plan for Journal Submission'
    doc.core_properties.subject = 'Evidence-based response to 14 experiment review issues'
    doc.core_properties.author = 'Research experiment review'
    doc.core_properties.keywords = 'GraphRAG, evaluation, reproducibility, journal revision'
    output = SOURCE.with_suffix('.docx')
    doc.save(output)
    # Bind this review to the inspected project snapshot, excluding credentials.
    selected = [ROOT/'README.md', ROOT/'requirements.txt']
    for folder, pattern in [('src','*.py'), ('config','*.yaml'), ('data/benchmark','*.csv'), ('data/benchmark','*.json'), ('results/final_evaluation','*_results.*'), ('results/final_evaluation_metrics','*.csv')]:
        selected.extend((ROOT/folder).rglob(pattern))
    selected.extend([ROOT/'results/statistical_tests.csv', ROOT/'results/scalability.csv', ROOT/'results/graph_validation_report.json', ROOT/'results/environment_manifest.json', ROOT/'results/final_evaluation_preview_model_partial/vector_rag_results.jsonl'])
    hashes = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(set(selected)) if p.is_file()}
    (HERE/'reviewed_artifact_hashes.json').write_text(json.dumps({'review_date':'2026-10-03','hash_algorithm':'SHA-256','artifacts':hashes},indent=2),encoding='utf-8')
    # Verify the generated package can be reopened and has all major sections.
    check = Document(output)
    assert len([p for p in check.paragraphs if p.style.name == 'Heading 1']) == 24
    assert len(check.tables) == source.count('\n|---')
    print(json.dumps({'word_document':str(output),'words_markdown':len(source.split()),'tables':len(check.tables),'major_sections':24,'snapshot_files_hashed':len(hashes),'bytes':output.stat().st_size},indent=2))


if __name__ == '__main__':
    build()
